"""Build a source-linked workbook. Recalculate its formulas with LibreOffice after export."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[2];DOCS=ROOT/'docs/rncp/bloc-03'
def read(name):return json.loads((DOCS/'donnees'/name).read_text())
facts=read('projet-reel.json');linear=read('linear-2026-09-14.json');b1=read('reference-bloc-01.json')
w=Workbook();w.remove(w.active)
def sheet(name,headers,rows,widths):
    s=w.create_sheet(name);s.append(headers)
    for row in rows:s.append(row)
    for c in s[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='12332D')
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='Arial',size=11,color='12332D');c.alignment=Alignment(vertical='top',wrap_text=True)
            if c.row%2==0:c.fill=PatternFill('solid',fgColor='EBEFE6')
        s.row_dimensions[row[0].row].height=44
    for i,width in enumerate(widths,1):s.column_dimensions[get_column_letter(i)].width=width
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    s.sheet_properties.pageSetUpPr.fitToPage=True
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    s.print_title_rows='1:1';s.print_options.horizontalCentered=True
    return s
sheet('Lecture',['Information','Portée'],[
['Projet','Spity, projet solo sur un an déclaré par Dorian. Mois exacts non fournis.'],
['Chronologie','Dates Git : étapes enregistrées, pas relevé exhaustif d’activité.'],
['Linear','Relevé du 14/09/2026, non actualisé. Décompte non pondéré.'],
['Cadrage','Charges et montants estimés du diaporama B1 ; aucun consommé réel.'],
['Navigation','Mesures locales du 15/09, un passage par parcours ; aucune mesure de production.'],
['Retours','Personas et retours simulés pour l’exercice ; pas d’entretiens réalisés.'],
['Formules','Recalculées par LibreOffice ; les sources sont conservées dans le kit.'],
],[24,110])
sheet('Chronologie',['Date Git','Étape','Révision','Source'],[[m['date'],m['title'],m['revision'],'preuves/historique-git.txt'] for m in facts['milestones']],[17,48,45,40])
sheet('Linear',['ID','Statut','Priorité','Intitulé','URL'],[[i['id'],i['status'],i['priority'],i['title'],i['url']] for i in linear['issues']],[14,18,17,72,90])
last=len(linear['issues'])+1
s=sheet('Indicateurs',['Indicateur','Valeur','Unité / limite'],[
['Tickets observés',f'=COUNTA(Linear!A2:A{last})','Relevé du 14/09'],
['Annulés',f'=COUNTIF(Linear!B2:B{last},"Canceled")','Exclus du dénominateur'],
['Non annulés','=B2-B3','Tickets, sans pondération'],
['Terminés',f'=COUNTIF(Linear!B2:B{last},"Done")','Statut historique'],
['Part terminée','=IF(B4=0,"",B5/B4)','Décompte ; ne mesure pas la livraison du produit'],
['Médiane avant','=MEDIAN(Navigation!B2:B6)','ms ; protocole local'],
['Médiane après','=MEDIAN(Navigation!C2:C6)','ms ; protocole local'],
],[34,22,85]);s['B6'].number_format='0.0%'
rows=[[l['title'],l['personDays'],b1['dailyRateEUR'],f'=B{i}*C{i}','Estimation B1'] for i,l in enumerate(b1['lots'],2)]
rows += [['Total développement','=SUM(B2:B10)',None,'=SUM(D2:D10)','Prévision, pas dépense réelle']]
rows += [[b['name'],None,None,b['amountEUR'],'Estimation B1'] for b in b1['budgetRows'][1:]]
rows += [['TOTAL B1',None,None,'=SUM(D11:D16)','Aucune convention HT ajoutée'],['Temps réellement consommé',None,None,None,'Non fourni : cellule vide, pas zéro'],['Coûts réellement consommés',None,None,None,'Non fournis : aucun écart calculé']]
s=sheet('Cadrage',['Lot / poste','j-h estimés','Taux €/j-h','Montant €','Nature'],rows,[45,18,18,20,60])
for row in s.iter_rows(min_row=2):row[3].number_format='#,##0.00'
p=facts['performance'];sheet('Navigation',['Parcours','Avant (ms)','Après (ms)','Écart (ms)','Protocole'],[[r,b,a,f'=C{i}-B{i}',p['scope']] for i,(r,b,a) in enumerate(zip(p['routes'],p['beforeMs'],p['afterMs']),2)],[24,18,18,18,100])
sheet('Retours SIMULES',['ID','Persona fictif','Retour inventé','Réponse pédagogique','Critère','Réalisation réelle'],[[r['id'],r['persona'],r['feedback'],r['response'],r['criterion'],r['implementationEvidence']] for r in facts['simulatedFeedback']],[16,34,65,65,60,25])
# Optional pedagogical extension must always remain separate from observations.
scenario=DOCS/'donnees/mise-en-situation.json'
if scenario.exists():
    sim=json.loads(scenario.read_text())
    if 'workbookRows' in sim:sheet('Suivi SIMULE',['Indicateur','Valeur','Nature'],sim['workbookRows'],[45,24,100])
output=ROOT/'tmp/bloc3/workbook-source/pilotage-spity.xlsx';output.parent.mkdir(parents=True,exist_ok=True)
w.save(output);print(output)
