"""Validate current Bloc 3 sources/exports and optionally assemble the offline kit.

Checks documentary consistency only: no application recipe or jury evaluation.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import zipfile
from urllib.parse import unquote
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[2];DOCS=ROOT/'docs/rncp/bloc-03';OUT=ROOT/'output/bloc-03'
PPTX=OUT/'spity-bloc-3-soutenance-v18.pptx';DECK=OUT/'spity-bloc-3-soutenance-v18.pdf'
PDF=OUT/'dossier-bloc-03-spity.pdf';XLSX=OUT/'pilotage-spity.xlsx';ZIP=OUT/'kit-soutenance-spity.zip'
def read(p):return json.loads(p.read_text())
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def git(*args):return subprocess.check_output(['git',*args],cwd=ROOT,text=True).strip()
def main(package):
    slides=read(DOCS/'donnees/support-oral.json');facts=read(DOCS/'donnees/projet-reel.json');sim=read(DOCS/'donnees/mise-en-situation.json')
    b1=read(DOCS/'donnees/reference-bloc-01.json');linear=read(DOCS/'donnees/linear-2026-09-14.json')
    coverage=read(DOCS/'donnees/couverture-criteres.json');sources=read(DOCS/'donnees/sources-evaluation.json')
    assert len(slides)==26 and [s['number'] for s in slides]==list(range(1,27))
    assert sum(s['minutes'] for s in slides)==30
    assert sum(s['minutes'] for s in slides if s['demo'])==6
    assert all(s['minutes']==0 for s in slides[23:])
    assert facts['team']['confirmedComposition']=='Projet solo : Dorian Joly'
    assert facts['duration']['label']=='1 an' and facts['actualHours'] is None and facts['actualCostEUR'] is None
    for src in sources:assert digest(ROOT/src['file'])==src['sha256']
    assert digest(ROOT/b1['source'])==b1['sourceSHA256']
    for m in facts['milestones']:
        assert git('show','-s','--format=%as',m['revision'])==m['date']
        assert subprocess.run(['git','merge-base','--is-ancestor',m['revision'],'872db19'],cwd=ROOT).returncode==0
    counts=Counter(i['status'] for i in linear['issues'])
    assert counts=={'Done':12,'Todo':4,'In Progress':1,'Backlog':7,'Canceled':1}
    assert sum(l['personDays'] for l in b1['lots'])==82
    assert sum(r['amountEUR'] for r in b1['budgetRows'])==44250
    t=sim['tracking']
    assert t['consumedPersonDays']+t['remainingPersonDays']==t['forecastPersonDays']==86
    assert t['forecastPersonDays']*t['dailyRateEUR']+t['fixedEstimatedEUR']==t['forecastEUR']==46050
    assert t['forecastEUR']-t['baselineEUR']==t['varianceEUR']==1800
    assert len(sim['planning'])==9 and sum(r['personDays'] for r in sim['planning'])==82
    for r in sim['planning']:assert 1<=r['startMonth']<=r['endMonth']<=12
    assert all(r['assignedHours']==4 and r['availableHours']==5 for r in sim['review']['roles'])
    assert len(coverage)==40
    competencies=['C.3.1','C3.2.1','C3.2.2','C3.3.1','C3.3.2','C3.4.1','C3.4.2']
    assert list(dict.fromkeys(c['competency'] for c in coverage))==competencies
    assert len({c['id'] for c in coverage})==40
    for c in coverage:
        assert all(1<=n<=26 for n in c['slides'])
        assert all(list((DOCS/'annexes').glob(a+'_*.md')) for a in c['annexes'])
        assert c['criterion'] in (DOCS/'CONTROLE_COMPLETUDE.md').read_text()
        assert any(c['competency'] in slides[n-1]['section'] for n in c['slides'])
    sections=list(dict.fromkeys(s['section'].split(' /')[0] for s in slides[3:23]))
    assert sections==competencies
    presentation=Presentation(PPTX);assert len(presentation.slides)==26
    charts=[]
    for info,slide in zip(slides,presentation.slides):
        assert info['notes']==info['script'], 'Les notes natives doivent contenir seulement le texte à dire'
        assert re.search(r'\bje\b|\bj[’\x27]|\bmon\b|\bmes\b',info['script'],re.I)
        assert info['notes']==slide.notes_slide.notes_text_frame.text
        assert info['script'] in (DOCS/'GUIDE_ORAL.md').read_text()
        text='\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert not re.search(r'15 jours|quinze jours|J1[0-5]\b|Camille|4270|4 270',text,re.I)
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=presentation.slide_width+10
            assert shape.top+shape.height<=presentation.slide_height+10
            if shape.has_chart:charts.append((info['number'],shape.chart))
    assert [n for n,c in charts]==[5,7,11]
    assert list(charts[0][1].series[0].values)==[r['startMonth']-1 for r in sim['planning']]
    assert list(charts[0][1].series[1].values)==[r['endMonth']-r['startMonth']+1 for r in sim['planning']]
    assert list(charts[1][1].series[0].values)==[12,1,4,7]
    perf=facts['performance']
    assert list(charts[2][1].series[0].values)==[v/1000 for v in perf['beforeMs']]
    assert list(charts[2][1].series[1].values)==[v/1000 for v in perf['afterMs']]
    with zipfile.ZipFile(PPTX) as z:
        assert z.testzip() is None
        assert len([n for n in z.namelist() if '/embeddings/' in n and n.endswith('.xlsx')])==3
    deck=PdfReader(DECK);dossier=PdfReader(PDF)
    assert len(deck.pages)==26 and len(dossier.pages)>10
    assert all(len(p.extract_text().strip())>70 for p in deck.pages)
    for n,page in enumerate(deck.pages,1):
        projected=page.extract_text()
        assert not re.search(r'\bminutes?\b|\bmin\b|\b\d{1,2}:\d{2}\b',projected,re.I), f'Chronométrage projeté : {n}'
        assert not re.search(r'repère privé|à montrer|ne pas lire|PDF fournis|acquisition présumée|ne pas déclarer',projected,re.I), f'Consigne interne projetée : {n}'
    assert all(len(p.extract_text().strip())>70 for p in dossier.pages)
    for n in [6,8,9,12,13,14,15,16,17,18,26]:
        content=deck.pages[n-1].extract_text().casefold()
        assert any(term in content for term in ['simul','mise en situation','pédagog','hypoth','proposé']),n
    wb=load_workbook(XLSX,data_only=False);cached=load_workbook(XLSX,data_only=True)
    assert wb.sheetnames==['Lecture','Chronologie','Linear','Indicateurs','Cadrage','Navigation','Retours SIMULES','Suivi SIMULE']
    expected={'Indicateurs':{'B2':25,'B3':1,'B4':24,'B5':12,'B6':.5,'B7':10220,'B8':748},'Cadrage':{'B11':82,'D11':36900,'D17':44250},'Suivi SIMULE':{'B6':86,'B9':44250,'B10':46050,'B11':1800,'B15':1,'B18':1.2}}
    for name,values in expected.items():
        for cell,value in values.items():assert cached[name][cell].value==value,(name,cell,cached[name][cell].value,value)
    for row in [18,19]:assert cached['Cadrage'][f'D{row}'].value is None
    formula_count=0
    for s in wb:
        for row in s:
            for c in row:
                assert c.data_type!='e'
                if c.data_type=='f':
                    formula_count+=1
                    assert cached[s.title][c.coordinate].value is not None
                    assert cached[s.title][c.coordinate].data_type!='e'
    visual=read(DOCS/'preuves/controle-visuel.json')
    for key,file in [('presentation',PPTX),('dossier',PDF)]:assert visual[key]['sha256']==digest(file)
    assert visual['presentation']['pdfSHA256']==digest(DECK)
    assert visual['presentation']['slidesReviewed']==list(range(1,27))
    assert visual['dossier']['pagesReviewed']==list(range(1,len(dossier.pages)+1))
    # Check actual destinations of all Markdown links in current documents and tool guide.
    links=0
    for p in [*DOCS.rglob('*.md'),ROOT/'tools/bloc3/README.md',ROOT/'output/README.md']:
        for target in re.findall(r'\[[^\]]*\]\(([^)]+)\)',p.read_text()):
            target=unquote(target.split('#')[0].strip('<>'))
            if not target or re.match(r'^[a-z]+:',target):continue
            dest=(p.parent/target).resolve()
            assert dest.exists() or (package and dest==ZIP),(p,target)
            links+=1
    current_tree=git('rev-parse','HEAD:spity');evidence=read(DOCS/'preuves/verification.json')
    assert git('rev-parse',f"{evidence['gitRevision']}:spity")==evidence['applicationTree']
    changes=git('diff','--name-only','HEAD','--','spity')
    report={'version':'v18','date':'2026-09-16','project':'Solo sur un an déclaré','competencies':competencies,'mandatoryCompetencies':['C.3.1','C3.2.1','C3.4.2'],'criteriaMapped':len(coverage),'slides':26,'presentationMinutes':30,'demonstrationMinutes':6,'nativeCharts':3,'pdfPages':len(dossier.pages),'workbookSheets':len(wb.sheetnames),'workbookFormulas':formula_count,'localLinksChecked':links,'applicationTree':current_tree,'verifiedHistoricalApplicationTree':evidence['applicationTree'],'applicationMatchesDatedEvidence':current_tree==evidence['applicationTree'] and not changes,'applicationWorkingTreeClean':not changes,'scope':'Cohérence documentaire, couverture des critères, formules, sources et exports. Réel et simulation distingués. Aucune nouvelle recette complète ni acquisition de compétence déclarée.'}
    (DOCS/'preuves/controle-kit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    if package:
        included=[p for p in DOCS.rglob('*') if p.is_file() and p.name!='MANIFEST.sha256']
        included += [p for p in (ROOT/'tools/bloc3').glob('*') if p.is_file()]
        included += [PDF,PPTX,DECK,XLSX,ROOT/b1['source'],ROOT/'docs/audits/2026-09-15-navigation-performances.md']
        included += [ROOT/src['file'] for src in sources]
        included += [ROOT/'spity/public/images/brand'/name for name in ['escalade-falaise-coucher-soleil.jpeg','escalade-falaise-gros-plan.jpeg']]
        included=sorted(set(included))
        manifest=DOCS/'preuves/MANIFEST.sha256';manifest.write_text(''.join(f'{digest(p)}  {p.relative_to(ROOT).as_posix()}\n' for p in included))
        included.append(manifest)
        with zipfile.ZipFile(ZIP,'w',zipfile.ZIP_DEFLATED) as z:
            for p in included:z.write(p,p.relative_to(ROOT).as_posix())
            z.writestr('LIRE_EN_PREMIER.txt','SPITY — BLOC 3 — v18\n\nOuvrir output/bloc-03/spity-bloc-3-soutenance-v18.pdf ou .pptx.\nSommaire, projet solo sur un an, puis sept compétences dans l’ordre.\nLes mises en situation pédagogiques sont signalées ; elles ne sont pas des expériences vécues.\nGuide : docs/rncp/bloc-03/GUIDE_ORAL.md\nCouverture : docs/rncp/bloc-03/CONTROLE_COMPLETUDE.md\nLa démonstration nécessite le dépôt applicatif complet et ses dépendances.\nCe kit ne contient pas de base, secrets ou node_modules. Aucun dépôt officiel ni accord client déclaré.\n')
        with zipfile.ZipFile(ZIP) as z:
            assert z.testzip() is None
            for p in included:assert hashlib.sha256(z.read(p.relative_to(ROOT).as_posix())).hexdigest()==digest(p)
        report['zipSHA256']=digest(ZIP)
    print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--package',action='store_true');main(parser.parse_args().package)
