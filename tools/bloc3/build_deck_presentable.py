"""Build the editable v17 deck from the solo project narrative and dated evidence.

Run from any directory: python tools/bloc3/build_deck_presentable.py
Dependencies: tools/bloc3/requirements-presentation.txt.
Rendering and visual review are separate; this generator never records a review.
"""
from pathlib import Path
import json
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / 'docs/rncp/bloc-03'
OUT = ROOT / 'output/bloc-03/spity-bloc-3-30-minutes-visuel-v17.pptx'
SOURCE = json.loads((DOCS / 'donnees/support-oral.json').read_text())
FACTS = json.loads((DOCS / 'donnees/projet-reel.json').read_text())
B1 = json.loads((DOCS / 'donnees/reference-bloc-01.json').read_text())
LINEAR = json.loads((DOCS / 'donnees/linear-2026-09-14.json').read_text())
SIM = json.loads((DOCS / 'donnees/mise-en-situation.json').read_text())
INK, GREEN, LIME, PAPER = '12332D', '286957', 'D8F28B', 'F6F5EF'
MUTED, LINE, WHITE, ORANGE = '5C6C65', 'DDE3D9', 'FFFFFF', 'B64E2E'
PALE, DARKCARD = 'EBEFE6', '1C453B'
FONT = 'Arial'
P = Presentation()
P.slide_width, P.slide_height = Inches(16), Inches(9)
P.core_properties.title = 'Spity — Piloter le projet | Bloc 3'
P.core_properties.subject = 'Soutenance RNCP39583 — 30 minutes, démonstration incluse'
P.core_properties.author = 'Dorian Joly'
P.core_properties.language = 'fr-FR'


def color(value):
    return RGBColor.from_string(value)


def rect(s, x, y, w, h, fill, radius=False, stroke=None):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = .04
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if stroke:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(.8)
    else:
        shape.line.fill.background()
    return shape


def txt(s, text, x, y, w, h, size=22, bold=False, fill=INK, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(text).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name, p.font.size, p.font.bold = FONT, Pt(size), bold
        p.font.color.rgb = color(fill)
        p.alignment = align
        p.space_after = Pt(4)
    return box


def line(s, x1, y1, x2, y2, fill=LINE, width=1):
    shape = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color(fill)
    shape.line.width = Pt(width)
    shape._element.spPr.append(OxmlElement('a:effectLst'))


def pill(s, text, x, y, w, dark=False):
    rect(s, x, y, w, .36, DARKCARD if dark else PALE, True)
    txt(s, text, x+.13, y+.065, w-.26, .24, 10, True, LIME if dark else GREEN)


def picture(s, relative, x, y, w, h, cover=False):
    path = ROOT / relative
    with Image.open(path) as im:
        ratio = im.width / im.height
    if cover:
        pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        target = w/h
        if ratio > target:
            pic.crop_left = pic.crop_right = (1-target/ratio)/2
        else:
            pic.crop_top = pic.crop_bottom = (1-ratio/target)/2
    else:
        iw, ih = (w, w/ratio) if ratio > w/h else (h*ratio, h)
        pic = s.shapes.add_picture(str(path), Inches(x+(w-iw)/2), Inches(y+(h-ih)/2), width=Inches(iw), height=Inches(ih))
    pic._element.nvPicPr.cNvPr.set('descr', path.stem)
    return pic


def base(n, title, section, subtitle='', dark=False):
    s = P.slides.add_slide(P.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(INK if dark else PAPER)
    txt(s, 'SPITY', .65, .35, 1.3, .28, 13, True, LIME if dark else INK)
    txt(s, section.upper(), 2.1, .38, 11, .25, 10, True, 'B9CABE' if dark else MUTED)
    if title:
        txt(s, title, .65, 1.02, 14.7, .78, 32, True, WHITE if dark else INK)
    if subtitle:
        txt(s, subtitle, .68, 1.87, 14.5, .65, 17, False, 'C7D6CB' if dark else MUTED)
    line(s, .65, 8.45, 15.35, 8.45, DARKCARD if dark else LINE)
    txt(s, 'DORIAN JOLY  /  RNCP39583 · BLOC 3', .65, 8.61, 7, .2, 9, False, 'B9CABE' if dark else MUTED)
    txt(s, f'{n:02d}', 14.6, 8.54, .75, .32, 16, True, LIME if dark else GREEN, PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = SOURCE[n-1]['notes']
    return s


def callout(s, text, y=7.65, dark=False):
    rect(s, .65, y, .055, .42, LIME if dark else GREEN)
    txt(s, text, .88, y+.025, 14.15, .55, 18, True, LIME if dark else GREEN)


def stat(s, value, label, x, y, w=3.5, dark=False, detail=None):
    txt(s, value, x, y, w, .95, 56, True, LIME if dark else GREEN)
    txt(s, label, x, y+1.02, w, .6, 20, True, WHITE if dark else INK)
    if detail:
        txt(s, detail, x, y+1.82, w, 1.1, 16, False, 'C7D6CB' if dark else MUTED)


def table(s, rows, widths, x=.65, y=2.65, h=4.6, size=17):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(sum(widths)), Inches(h)).table
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(h/len(rows))
        for ci, value in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = str(value)
            cell.margin_left = Inches(.15)
            cell.margin_right = Inches(.12)
            cell.margin_top = Inches(.07)
            cell.margin_bottom = Inches(.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = color(INK if ri == 0 else WHITE if ri%2 else PALE)
            for p in cell.text_frame.paragraphs:
                p.font.name, p.font.size = FONT, Pt(size)
                p.font.bold = ri == 0 or ci == 0
                p.font.color.rgb = color(WHITE if ri == 0 else INK)
            props = cell._tc.get_or_add_tcPr()
            for tag in ['lnL', 'lnR', 'lnT', 'lnB']:
                edge = OxmlElement('a:'+tag)
                edge.append(OxmlElement('a:noFill'))
                props.append(edge)
    return t


def chart(s, categories, series, x, y, w, h, maximum, dark=False, stacked=False,
          fmt='0', colors=None, labels=True, legend=False, minimum=0, reverse=True):
    data = CategoryChartData()
    data.categories = categories
    for name, vals in series:
        data.add_series(name, vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED,
                           Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    ch.chart_style = 10
    ch.has_title = False
    ch.font.name = FONT
    ch.font.size = Pt(16)
    ch.font.color.rgb = color(WHITE if dark else INK)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.name, ch.legend.font.size = FONT, Pt(13)
    ch.category_axis.reverse_order = reverse
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ch.category_axis.tick_labels.font.size = Pt(16)
    ch.value_axis.minimum_scale = minimum
    ch.value_axis.maximum_scale = maximum
    ch.value_axis.tick_labels.font.size = Pt(12)
    ch.value_axis.tick_labels.number_format = fmt
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = color(DARKCARD if dark else LINE)
    ch.value_axis.major_gridlines.format.line.width = Pt(.6)
    for axis in [ch.category_axis, ch.value_axis]:
        axis.format.line.fill.background()
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.color.rgb = color('C7D6CB' if dark else MUTED)
    plot = ch.plots[0]
    plot.gap_width = 65
    plot.has_data_labels = labels
    if labels:
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END if not stacked else XL_LABEL_POSITION.CENTER
        plot.data_labels.font.name, plot.data_labels.font.size = FONT, Pt(17)
        plot.data_labels.font.bold = True
        plot.data_labels.font.color.rgb = color(WHITE if dark else INK)
        plot.data_labels.number_format = fmt
    for i, ser in enumerate(ch.series):
        fill = (colors or [GREEN, 'BBCBAE'])[i % len(colors or [GREEN, 'BBCBAE'])]
        if fill is None:
            ser.format.fill.background()
        else:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = color(fill)
        ser.format.line.fill.background()
    # Prevent theme-dependent chart-area backgrounds in PowerPoint / LibreOffice.
    for parent in [ch._chartSpace, ch._chartSpace.chart.plotArea]:
        props = OxmlElement('c:spPr')
        props.append(OxmlElement('a:noFill'))
        border = OxmlElement('a:ln')
        border.append(OxmlElement('a:noFill'))
        props.append(border)
        parent.append(props)
    return ch


def points(ch, fills):
    for point, fill in zip(ch.series[0].points, fills):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color(fill)
        point.format.line.fill.background()



for item in SOURCE:
    n, layout, content = item['number'], item['layout'], item['content']
    dark = item['dark']
    fg, secondary = (WHITE, 'C7D6CB') if dark else (INK, MUTED)
    s = base(n, '' if layout in ['cover', 'demo'] else item['title'], item['section'],
             '' if layout in ['cover', 'demo'] else item['subtitle'], dark)
    if layout in ['cover', 'demo']:
        photo = 'escalade-falaise-coucher-soleil.jpeg' if layout == 'cover' else 'escalade-falaise-gros-plan.jpeg'
        picture(s, f'spity/public/images/brand/{photo}', 8.85, 0, 7.15, 8.4, True)
        pill(s, 'PROJET SOLO · 1 AN' if layout == 'cover' else 'DÉMONSTRATION EN DIRECT', .75, 1.6, 3.6, True)
        txt(s, 'Une année\npour construire\nSpity.' if layout == 'cover' else 'Passer du suivi\nà l’usage.', .7, 2.5, 8, 3.1, 55, True, WHITE)
        txt(s, item['subtitle'], .75, 6, 7.35, 1.0, 24, False, 'C7D6CB')
        txt(s, 'Dorian Joly' if layout == 'cover' else '6 minutes', .75, 7.45, 7, .55, 25, True, LIME)
    elif layout == 'agenda':
        stat(s, '30', 'minutes pour présenter', .75, 3, 4, True, '24 min d’explications\n6 min de démonstration')
        for i,(num,label,timing) in enumerate(content['items']):
            y=2.67+i*.65
            txt(s,num,5.35,y,.7,.5,24,True,LIME)
            txt(s,label,6.15,y,6.1,.5,21,True,WHITE)
            txt(s,timing,12.6,y+.09,2.6,.35,14,False,'C7D6CB',PP_ALIGN.RIGHT)
            line(s,6.15,y+.55,15.15,y+.55,DARKCARD)
    elif layout == 'product':
        for i,(label,heading,detail) in enumerate(content['cards']):
            y=2.8+i*2.05
            pill(s,label,.7,y,1.8)
            txt(s,heading,.7,y+.55,5.7,.6,25,True)
            txt(s,detail,.7,y+1.17,5.7,.9,18,False,MUTED)
        rect(s,6.85,2.72,8.5,4.82,INK,True)
        picture(s,'docs/rncp/bloc-03/preuves/captures/lieux-2026-09-15.png',7.03,2.88,8.14,4.46)
    elif layout == 'timeline':
        stat(s,'1 an','projet solo',.8,2.9,3.0,False,'Période déclarée\npar Dorian Joly')
        for i,(date,title,detail) in enumerate(content['items']):
            y=2.65+i*.91
            pill(s,date,4.4,y,1.6)
            txt(s,title,6.28,y-.02,8.5,.45,23,True)
            txt(s,detail,6.28,y+.46,8.5,.32,16,False,MUTED)
            if i<4:line(s,6.28,y+.81,15.1,y+.81)
    elif layout == 'annual':
        seq=SIM['planning']
        chart(s,[r['lot'] for r in seq],
              [('Début',[r['startMonth']-1 for r in seq]),
               ('Durée',[r['endMonth']-r['startMonth']+1 for r in seq])],
              .6,2.65,11.0,4.5,12,stacked=True,colors=[None,GREEN],labels=False,fmt='0')
        txt(s,'MOIS RELATIFS · PLANNING RÉVISÉ DU CAS',.9,7.16,10.7,.27,11,True,MUTED)
        stat(s,'82','j-h estimés B1',11.9,3.03,3.5,False,'Dépendances :\naccès → profils\névénement → inscription')
    elif layout == 'effort':
        seq=B1['lots']
        chart(s,[r['shortTitle'] for r in seq],[('Charge estimée',[r['personDays'] for r in seq])],.65,2.65,10.65,4.62,14,fmt='0" j-h"')
        rect(s,11.6,2.9,3.75,4.3,PALE,True)
        stat(s,str(B1['personDays']),'jours-personne',11.9,3.3,3.1,False,'Estimation initiale\nAucun relevé réel\nde temps fourni')
    elif layout == 'linear':
        statuses=['Done','In Progress','Todo','Backlog']
        counts=[sum(i['status']==status for i in LINEAR['issues']) for status in statuses]
        assert sum(counts)==24, counts
        chart(s,['Terminé','En cours','À faire','Backlog'],[('Tickets',counts)],.65,2.9,10.65,4.0,14)
        stat(s,'12 / 24','tickets terminés',11.65,3.1,3.7,False,'14 septembre 2026\n1 annulé exclu\nTailles non pondérées')
    elif layout == 'performance':
        data=FACTS['performance']
        chart(s,data['routes'],[('Avant',[v/1000 for v in data['beforeMs']]),('Après',[v/1000 for v in data['afterMs']])],.65,2.7,10.7,4.6,13,True,fmt='0.00" s"',colors=['86A18C',LIME],legend=True)
        txt(s,'MÉDIANE OBSERVÉE',11.6,3.05,3.8,.4,13,True,'C7D6CB')
        txt(s,'10,22 s',11.6,3.76,3.6,.8,40,True,WHITE)
        txt(s,'↓',11.65,4.75,2,.65,30,True,LIME)
        txt(s,'0,748 s',11.6,5.48,3.6,.8,40,True,LIME)
        txt(s,'5 parcours locaux',11.6,6.6,3.7,.42,17,False,'C7D6CB')
    elif layout == 'cards':
        for i,(tag,title,detail) in enumerate(content['cards']):
            x=.65+i*5.02
            rect(s,x,2.85,4.65,4.3,DARKCARD if dark else WHITE,True)
            txt(s,tag,x+.26,3.15,4.1,.57,13,True,LIME if dark else GREEN)
            txt(s,title,x+.26,4.0,4.1,1.06,28,True,fg)
            txt(s,detail,x+.26,5.4,4.08,1.4,20,False,secondary)
    elif layout == 'management':
        for i,(title,verb,detail) in enumerate(content['cards']):
            x=.65+(i%2)*7.52; y=2.75+(i//2)*2.04
            rect(s,x,y,7.18,1.81,DARKCARD,True)
            txt(s,title.upper(),x+.25,y+.17,6.7,.3,12,True,LIME)
            txt(s,verb,x+.25,y+.6,2.6,.85,26,True,WHITE)
            txt(s,detail,x+3.08,y+.59,3.78,1.0,18,False,'C7D6CB')
    elif layout in ['table','dashboard','skills','training']:
        rows=content['rows']
        widths=[3.0,5.45,6.25] if len(rows[0])==3 else [3.5,11.2]
        table(s,rows,widths,y=2.65,h=4.66,size=17 if len(rows)>6 else 19)
    elif layout == 'screenshot':
        rect(s,.65,2.5,9.2,5.65,INK,True)
        picture(s,f"docs/rncp/bloc-03/preuves/captures/{content['image']}-2026-09-15.png",.79,2.6,8.92,5.44)
        for i,label in enumerate(content['steps']):
            y=3.15+i*1.34
            pill(s,f'0{i+1}',10.7,y,1.0)
            txt(s,label,10.7,y+.48,4.4,.55,28,True)
        txt(s,'Capture de secours · 15/09/2026',10.7,7.83,4.5,.3,11,False,MUTED)
    else:
        raise ValueError(layout)
    if content.get('callout'):
        # Keep the source/limitation readable without overflowing the footer.
        rect(s,.65,7.65,.055,.48,LIME if dark else GREEN)
        txt(s,content['callout'],.88,7.65,14.25,.67,17,True,LIME if dark else GREEN)

assert len(P.slides) == len(SOURCE) == 26
assert sum(s['minutes'] for s in SOURCE) == 30
assert sum(s['minutes'] for s in SOURCE if s['demo']) == 6
for n, slide in enumerate(P.slides, 1):
    assert slide.notes_slide.notes_text_frame.text == SOURCE[n-1]['notes']
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, (n, shape.name)
        assert shape.left+shape.width <= P.slide_width+10, (n, shape.name, 'right')
        assert shape.top+shape.height <= P.slide_height+10, (n, shape.name, 'bottom')
OUT.parent.mkdir(parents=True, exist_ok=True)
P.save(OUT)
print(OUT)
